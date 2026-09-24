"""Build a lossless content manifest and web-sized copies from NEWDESIGN sources."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parents[1]
PHOTO_ROOT = ROOT / "照片" / "家具"
ASSET_ROOT = OUT / "assets"

# Excel row order is the authority. Folder 2 uses an older "嶼 ISLE" label;
# retain the Excel public name until the client resolves that discrepancy.
WORKS = [
    ("ridge-table", "1稜 RIDGE", "未加logo.png", "tables"),
    ("landscape-table", "2嶼 ISLE", "原木桌.png", "tables"),
    ("muju", "1栓木桌 _木聚", "栓木餐桌椅與乾燥花.png", "tables"),
    ("musu", "2硬楓木長桌", "極簡長桌乾燥花.png", "tables"),
    ("muyo", "3胡桃木餐桌椅", "胡桃曲木餐桌椅.png", "tables"),
    ("muwi", "4胡桃曲木桌", "視覺.png", "tables"),
    ("muki", "5日式風格桌", "折界餐桌日式風格.png", "tables"),
    ("muni", "6 木日 MUNI", "木日 MUNI.png", "tables"),
    ("ripple", "漾 RIPPLE", "漾 RIPPLE｜拼木圓椅.png", "seating"),
    ("warm", "暖 WARM", "暖 WARM｜實木圓凳01.JPG", "seating"),
    ("walnut-island-bar", "胡桃木中島酒櫃", "胡桃木中島酒櫃.png", "islands"),
    ("cherry-island-table", "櫻桃木中島桌", "階梯櫻桃木中島.png", "islands"),
    ("grid", "格櫃｜GRID", "格｜GRID.png", "storage"),
    ("ridge-shelf", "嶺 RIDGE", "嶺 RIDGE｜實木層架 正面01.png", "storage"),
    ("arch", "築｜ARCH", "築｜ARCH ＊1櫃.png", "storage"),
    ("plain", "原 PLAIN", "原 PLAIN｜實木書架正面.png", "storage"),
    ("breeze", "風 BREEZE", "風 BREEZE｜實木百葉收納櫃.png", "storage"),
]


def web_copy(source: Path, target: Path, *, width: int = 1800) -> str:
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        image = ImageOps.exif_transpose(image)
        if image.width > width:
            image = image.resize((width, round(image.height * width / image.width)), Image.Resampling.LANCZOS)
        if image.mode not in ("RGB", "RGBA"):
            image = image.convert("RGB")
        image.save(target, "WEBP", quality=86, method=6)
    return "/assets/" + target.relative_to(ASSET_ROOT).as_posix()


def find_folder(part: str) -> Path:
    matches = [p for p in PHOTO_ROOT.rglob("*") if p.is_dir() and part in p.name]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one folder for {part}, found {matches}")
    return matches[0]


sheet = load_workbook(
    ROOT / "網頁分層圖示及文字" / "官網作品內頁版型" / "飛熊入夢_官網作品資料總表.xlsx",
    read_only=True, data_only=True,
).worksheets[0]
rows = list(sheet.iter_rows(values_only=True))[1:]
assert len(rows) == len(WORKS), (len(rows), len(WORKS))

works = []
for row, (slug, folder_part, hero_name, category) in zip(rows, WORKS):
    folder = find_folder(folder_part)
    photos = sorted(p for p in folder.iterdir() if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"})
    heroes = [p for p in photos if p.name == hero_name]
    if len(heroes) != 1:
        raise RuntimeError(f"Hero not found: {slug}: {hero_name}")
    photos.remove(heroes[0])
    photos.insert(0, heroes[0])
    if slug == "ridge-table":
        photos = photos[:1]  # The other supplied versions contain text and logo baked into pixels.
    out_dir = ASSET_ROOT / "works" / slug
    for stale in out_dir.glob("*.webp"):
        if int(stale.stem) > len(photos):
            stale.unlink()
    gallery = [web_copy(photo, ASSET_ROOT / "works" / slug / f"{i:02d}.webp") for i, photo in enumerate(photos, 1)]
    works.append({
        "slug": slug, "series": row[0], "title": row[1], "english": row[2],
        "type": row[3], "material": row[4], "size": row[5], "tagline": row[6],
        "story": row[7], "craft": row[8], "custom": row[9], "note": row[10],
        "category": category, "image": gallery[0], "gallery": gallery,
        "source_folder": folder.relative_to(ROOT).as_posix(),
    })

tray_root = ROOT / "照片" / "生活木作" / "晨露圓境托盤_ Alba Canvas"
tray_files = sorted(p for p in tray_root.iterdir() if p.suffix.lower() in {".jpg", ".png"})
tray = [web_copy(p, ASSET_ROOT / "lifestyle" / f"tray-{i:02d}.webp") for i, p in enumerate(tray_files, 1)]
workshop = ROOT / "照片" / "工作照" / "IMG_1712.JPG"
web_copy(workshop, ASSET_ROOT / "brand" / "real-lecture.webp")
poster_root = ROOT / "網頁分層圖示及文字" / "木作學堂頁面" / "課程簡章_海報"
course_sources = [
    ("beginner", "木工基礎入門班", "1初階木工精華班.png", "6 小時", "3680", "初學者適合"),
    ("cnc", "CNC 數位木工", "3ＣＮＣ數位木工.png", "6 小時", "8800", "有基礎佳"),
    ("sharpening", "磨刀實戰班", "2木工磨刀實戰班.png", "6 小時", "3680", "工具保養"),
    ("open-studio", "自由創作會員", "4.會員自由創作.png", "彈性時段", "1800", "進階創作"),
]
courses = []
for slug, title, source_name, duration, price, level in course_sources:
    photo = web_copy(poster_root / source_name, ASSET_ROOT / "courses" / f"{slug}.webp", width=1200)
    courses.append({"slug": slug, "title": title, "duration": duration, "price": price, "level": level, "mode": "onsite", "image": photo})
for i, p in enumerate(sorted((ROOT / "照片" / "最新消息").glob("*")), 1):
    if p.suffix.lower() in {".jpg", ".png"}:
        web_copy(p, ASSET_ROOT / "journal" / f"journal-{i:02d}.webp")

manifest = {
    "works": works,
    "lifestyle": {"title": "晨露圓境托盤", "english": "Alba Canvas", "gallery": tray},
    "courses": courses,
    "source": "NEWDESIGN customer Excel and original photos",
}
(OUT / "data").mkdir(exist_ok=True)
(OUT / "data" / "content.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Wrote {len(works)} works, {sum(len(w['gallery']) for w in works)} furniture photos, {len(tray)} tray photos")
